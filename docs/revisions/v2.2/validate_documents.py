"""Validate the v2.2 document package; --write-index explicitly regenerates its index."""
from pathlib import Path
from functools import lru_cache
from zipfile import ZipFile
import argparse
import csv
import hashlib
import json
import re
from docx import Document

folder = Path(__file__).resolve().parent
root = folder.parents[2]
previous = folder.parent / 'v2.1'
source = next(previous.glob('*.docx'))
revised = next(folder.glob('*.docx'))
new_ids = {f'MAIL-{n}' for n in range(17, 23)}
new_cards = {'M4-15', 'M4-16', 'M4-17', 'M4-18', 'M4-19',
             'M5-11', 'M5-12', 'M6-11', 'M6-12', 'M6-13',
             'M8-15', 'M8-16', 'M8-17', 'M9-19'}
retired_cards = {'M10-01', 'M10-02'}


def requirements(path):
    result = {}
    for table in Document(path).tables:
        if len(table.columns) != 5:
            continue
        for row in table.rows[1:]:
            match = re.match(r'^([A-Z]+-\d{2})\b', row.cells[0].text)
            if match:
                assert match[1] not in result, ('duplicate requirement', match[1])
                result[match[1]] = [cell.text for cell in row.cells]
    return result


def cards(path):
    result = {}
    for line in path.read_text(encoding='utf-8-sig').splitlines():
        columns = [c.strip() for c in line.split('|')[1:-1]]
        if len(columns) == 6 and re.fullmatch(r'M\d+-\d+', columns[0]):
            assert columns[0] not in result, ('duplicate card', columns[0])
            result[columns[0]] = columns
    return result


def validate(write_index=False):
    before, after = requirements(source), requirements(revised)
    assert len(before) == 161 and len(after) == 167
    assert before.keys() <= after.keys() and after.keys() - before.keys() == new_ids
    assert all(before[k][4] == after[k][4] for k in before), 'existing stage changed'
    promoted = {k for k in before if before[k][3] != after[k][3]}
    assert promoted == {'MAIL-13', 'AI-12'}
    assert all(before[k][3] == 'P1' and after[k][3] == 'P0' for k in promoted)
    assert all(after[k][3:] == ['P0', '1'] for k in new_ids)
    changes = json.loads((folder / 'requirements-changes.json').read_text(encoding='utf-8'))
    actual_changes = {k for k in after if k not in before or after[k] != before[k]}
    assert len(changes) == 14 and {r['id'] for r in changes} == actual_changes
    for item in changes:
        assert item['before'] == before.get(item['id'])
        assert item['after'] == after[item['id']]

    doc = Document(revised)
    assert len(doc.inline_shapes) == 2
    with ZipFile(source) as old_zip, ZipFile(revised) as new_zip:
        old_images = {n: old_zip.read(n) for n in old_zip.namelist() if n.startswith('word/media/')}
        new_images = {n: new_zip.read(n) for n in new_zip.namelist() if n.startswith('word/media/')}
        assert old_images == new_images, 'original figure bytes changed'
    send = next(p.text for p in doc.paragraphs if '"operation_id"' in p.text and '"draft_id"' in p.text)
    assert {'operation_id', 'expected_version', 'payload_hash', 'approval_id'} <= json.loads(send).keys()
    body = '\n'.join(p.text for p in doc.paragraphs)
    assert '모든 대량 변경(신규·기존 메일 포함)' in body
    for number in range(26, 31):
        assert any(p.text.startswith(f'{number} ') for p in doc.paragraphs)
    assert any('mail.folders.create / rename / move / remove' in c.text
               for t in doc.tables for r in t.rows for c in r.cells)
    for checksum_file in [folder / 'source-files.sha256', previous / 'original-files.sha256']:
        for line in checksum_file.read_text(encoding='utf-8').splitlines():
            digest, name = line.split('  ', 1)
            assert hashlib.sha256((root / name).read_bytes()).hexdigest() == digest, name

    old_plan, plan = cards(previous / 'stage1-tasks.md'), cards(folder / 'stage1-tasks.md')
    assert len(old_plan) == 135 and len(plan) == 147
    assert plan.keys() - old_plan.keys() == new_cards
    assert old_plan.keys() - plan.keys() == retired_cards

    def dependencies(card):
        text, found = plan[card][3], set()
        for match in re.finditer(r'\bM(\d+) 나머지 전부', text):
            found.update(k for k in plan if k.startswith('M' + match[1] + '-') and k != card)
        pattern = r'\bM(\d+)-(\d{2})(?:~(?:M(\d+)-)?(\d{2}))?'
        for match in re.finditer(pattern, text):
            start_m, start_n = int(match[1]), int(match[2])
            if match[4]:
                end_m, end_n = int(match[3] or match[1]), int(match[4])
                if start_m == end_m:
                    found.update(f'M{start_m}-{n:02}' for n in range(start_n, end_n + 1))
                else:
                    assert start_n == end_n, ('ambiguous dependency range', card)
                    found.update(f'M{m}-{start_n:02}' for m in range(start_m, end_m + 1))
            else:
                found.add(f'M{start_m}-{start_n:02}')
        assert card not in found, (card, 'self-dependency')
        assert found <= plan.keys(), (card, 'unknown dependency', found - plan.keys())
        return found

    edges = {k: dependencies(k) for k in plan}
    visiting, visited = set(), set()

    def visit(card):
        assert card not in visiting, ('dependency cycle', card)
        if card in visited:
            return
        visiting.add(card)
        for other in edges[card]:
            visit(other)
        visiting.remove(card)
        visited.add(card)

    for card in plan:
        visit(card)

    @lru_cache(None)
    def ancestors(card):
        result = set()
        for dep in edges[card]:
            result.add(dep)
            result.update(ancestors(dep))
        return result

    gate = ancestors('M9-12')
    assert new_cards <= gate, ('new P0 card outside Gate 1', new_cards - gate)
    assert not {'M5-07', 'M8-14', 'M9-13', 'M10-03', 'M10-04'} & gate
    assert {'M4-14', 'M5-08', 'M7-06', 'M8-10'} <= gate
    assert 'M6-13' in ancestors('M8-15')

    mapping = {}
    for card, columns in plan.items():
        for match in re.finditer(r'\b([A-Z]+)-(\d{2})(?:~(?:[A-Z]+-)?(\d{2}))?', columns[2]):
            for n in range(int(match[2]), int(match[3] or match[2]) + 1):
                rid = f'{match[1]}-{n:02}'
                assert rid in after, (card, rid)
                mapping.setdefault(rid, set()).add(card)
    p0 = [k for k, r in after.items() if r[3] == 'P0' and r[4].startswith('1')]
    assert len(p0) == 141
    assert all(mapping.get(k, set()) & gate for k in p0), [k for k in p0 if not mapping.get(k, set()) & gate]
    with (folder / 'card-status.csv').open(encoding='utf-8-sig', newline='') as stream:
        statuses = list(csv.DictReader(stream))
    assert {r['card_id'] for r in statuses} == set(plan) and len(statuses) == len(plan)
    assert all(r['status'] == 'not_started' and not r['reviewer'] and not r['approved_at_utc']
               and not r['evidence'] and not r['blocking_reason'] for r in statuses)
    assert all(r['title'] == plan[r['card_id']][1] for r in statuses)
    assert all(r['review_type'] == plan[r['card_id']][5] for r in statuses)
    assert all(r['dependencies'] == plan[r['card_id']][3] for r in statuses)

    index_path = folder / 'requirements-index.csv'
    fields = ['requirement_id', 'title', 'description', 'acceptance', 'priority', 'stage',
              'stage1_card_ids', 'verification_status']
    if write_index:
        with index_path.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(fields)
            for rid, record in after.items():
                writer.writerow([rid, *record, ','.join(sorted(mapping.get(rid, set()))), 'not_implemented'])
    with index_path.open(encoding='utf-8-sig', newline='') as stream:
        index = list(csv.DictReader(stream))
    assert len(index) == 167 and {r['requirement_id'] for r in index} == set(after)
    for item in index:
        rid = item['requirement_id']
        assert [item[k] for k in fields[1:6]] == after[rid], rid
        assert item['stage1_card_ids'] == ','.join(sorted(mapping.get(rid, set()))), rid
        assert item['verification_status'] == 'not_implemented', rid
    return {'requirements': len(after), 'new_requirements': sorted(new_ids),
            'promoted_to_P0': sorted(promoted), 'existing_stages_preserved': True,
            'figures_preserved': 2, 'active_cards': len(plan), 'new_cards': len(new_cards),
            'retired_cards': sorted(retired_cards), 'dependency_cycles': 0,
            'gate1_optional_P1_dependencies': 0, 'stage1_P0_mapped_to_gate1': len(p0),
            'original_and_v21_hashes_unchanged': True,
            'card_statuses': 'all not_started; no approval recorded',
            'word_sha256': hashlib.sha256(revised.read_bytes()).hexdigest()}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--write-index', action='store_true')
    args = parser.parse_args()
    print(json.dumps(validate(args.write_index), ensure_ascii=False, indent=2))
