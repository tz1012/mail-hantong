"""Read-only checks for the v2.1 document revision package. Requires python-docx."""
from pathlib import Path
from docx import Document
import csv,hashlib,re,json
from functools import lru_cache
from zipfile import ZipFile

folder=Path(__file__).resolve().parent
root=folder.parents[2]
original=next(root.glob('*.docx'))
revised=next(folder.glob('*.docx'))

def requirements(path):
    result={}
    for t in Document(path).tables:
        if len(t.columns)!=5:continue
        for r in t.rows[1:]:
            m=re.match(r'^([A-Z]+-\d{2})\b',r.cells[0].text)
            if m:
                assert m[1] not in result
                result[m[1]]=[c.text for c in r.cells]
    return result

before=requirements(original);after=requirements(revised)
assert len(before)==161 and before.keys()==after.keys()
assert all(before[k][3:]==after[k][3:] for k in before)
doc=Document(revised)
assert len(doc.inline_shapes)==2
with ZipFile(original) as old_zip, ZipFile(revised) as new_zip:
    old_images={n:old_zip.read(n) for n in old_zip.namelist() if n.startswith('word/media/')}
    new_images={n:new_zip.read(n) for n in new_zip.namelist() if n.startswith('word/media/')}
    assert old_images==new_images,'original figure bytes changed'
body='\n'.join(p.text for p in doc.paragraphs)
send=next(p.text for p in doc.paragraphs if '"operation_id"' in p.text and '"draft_id"' in p.text)
assert all(k in json.loads(send) for k in ['operation_id','expected_version','payload_hash','approval_id'])
for line in (folder/'original-files.sha256').read_text(encoding='utf-8').splitlines():
    digest,name=line.split('  ',1)
    assert hashlib.sha256((root/name).read_bytes()).hexdigest()==digest,name

def cards(path):
    result={}
    for line in path.read_text(encoding='utf-8-sig').splitlines():
        cols=[c.strip() for c in line.split('|')[1:-1]]
        if len(cols)==6 and re.fullmatch(r'M\d+-\d+',cols[0]):
            assert cols[0] not in result
            result[cols[0]]=cols
    return result

plan=cards(folder/'stage1-tasks.md')
assert plan.keys()==cards(root/'stage1-tasks_1.md').keys()
assert len(plan)==135

def dependencies(card):
    text=plan[card][3];found=set()
    for m in re.finditer(r'\bM(\d+) 나머지 전부',text):
        found.update(k for k in plan if k.startswith('M'+m[1]+'-') and k!=card)
    pattern=r'\bM(\d+)-(\d{2})(?:~(?:M(\d+)-)?(\d{2}))?'
    for m in re.finditer(pattern,text):
        start_m,start_n=int(m[1]),int(m[2])
        if m[4]:
            end_m=int(m[3] or m[1]);end_n=int(m[4])
            if start_m==end_m:found.update(f'M{start_m}-{n:02}' for n in range(start_n,end_n+1))
            else:
                assert start_n==end_n
                found.update(f'M{milestone}-{start_n:02}' for milestone in range(start_m,end_m+1))
        else:found.add(f'M{start_m}-{start_n:02}')
    assert card not in found,(card,'self-dependency')
    assert found<=plan.keys(),(card,'unknown dependency',found-plan.keys())
    return found

edges={k:dependencies(k) for k in plan}
visiting=set();visited=set()
def visit(k):
    assert k not in visiting,('cycle',k)
    if k in visited:return
    visiting.add(k)
    for other in edges[k]:visit(other)
    visiting.remove(k);visited.add(k)
for k in plan:visit(k)
@lru_cache(None)
def ancestors(k):
    result=set()
    for dep in edges[k]:result.add(dep);result.update(ancestors(dep))
    return result
gate=ancestors('M9-12')
assert not {'M5-07','M8-14','M9-13','M10-01','M10-02','M10-03','M10-04'} & gate
assert {'M4-14','M5-08','M7-06','M8-10'}<=gate

mapping={}
for cid,cols in plan.items():
    for m in re.finditer(r'\b([A-Z]+)-(\d{2})(?:~(?:[A-Z]+-)?(\d{2}))?',cols[2]):
        for n in range(int(m[2]),int(m[3] or m[2])+1):
            rid=f'{m[1]}-{n:02}'
            assert rid in after,(cid,rid)
            mapping.setdefault(rid,set()).add(cid)
p0=[k for k,r in after.items() if r[3]=='P0' and r[4].startswith('1')]
assert all(k in mapping for k in p0),[k for k in p0 if k not in mapping]
with (folder/'card-status.csv').open(encoding='utf-8-sig',newline='') as f:status=list(csv.DictReader(f))
assert {r['card_id'] for r in status}==set(plan) and len(status)==len(plan)
assert all(r['status']=='not_started' and not r['reviewer'] and not r['approved_at_utc'] for r in status)
assert all(r['dependencies']==plan[r['card_id']][3] for r in status)
index_path=folder/'requirements-index.csv'
if index_path.exists():
    with index_path.open(encoding='utf-8-sig',newline='') as f:index=list(csv.DictReader(f))
    assert len(index)==161 and {r['requirement_id'] for r in index}==set(after)
    for item in index:
        record=after[item['requirement_id']]
        assert [item[k] for k in ['title','description','acceptance','priority','stage']]==record
print(json.dumps({'requirements':len(after),'IDs_priorities_stages_preserved':True,'figures':2,'cards':len(plan),'dependency_cycles':0,'gate1_optional_P1_dependencies':0,'stage1_P0_requirements_mapped':len(p0),'original_hashes_unchanged':True,'card_statuses':'all not_started; no approval recorded'},ensure_ascii=False,indent=2))
